import io

import openpyxl
import pytest
from docx import Document

HTML = b"""<html><head><style>x{}</style><title>t</title></head>
<body><h1>Pliego</h1><p>Objeto del contrato: plataforma de datos.</p>
<script>ignore()</script><p>Segundo parrafo.</p></body></html>"""


@pytest.fixture
def html_bytes() -> bytes:
    return HTML


@pytest.fixture
def docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Pliego de prescripciones", level=1)
    doc.add_paragraph("Objeto: plataforma de datos y analítica.")
    doc.add_paragraph("Requisitos: integración de APIs.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Criterios"
    ws.append(["Criterio", "Puntos"])
    ws.append(["Precio", 40])
    ws.append(["Técnica", 60])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
